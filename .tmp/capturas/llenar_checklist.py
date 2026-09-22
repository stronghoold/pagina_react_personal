"""Llena la lista de chequeo del cuarto avance (React + FastAPI).

- Marca el estado de los 26 requerimientos.
- Escribe las observaciones del aprendiz.
- Inserta una captura de evidencia por requerimiento (todas en marco Windows).

Entrada : C:\\Users\\simon\\Downloads\\Lista_Chequeo_Cuarto_Avance_REACT_FASTAPI (1).xlsx
Salida  : Lista_Chequeo_Cuarto_Avance_REACT_FASTAPI_LLENADA.xlsx (raiz del proyecto)
"""
import os

import openpyxl
from openpyxl.drawing.image import Image as XLImage

ROOT = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
TEMPLATE = r"C:\Users\simon\Downloads\Lista_Chequeo_Cuarto_Avance_REACT_FASTAPI (1).xlsx"
CAPTURAS = os.path.join(ROOT, "docs", "capturas-react")
SALIDA = os.path.join(ROOT, "Lista_Chequeo_Cuarto_Avance_REACT_FASTAPI_LLENADA.xlsx")

ESTADO = "Cumple"

# fila -> captura de evidencia
EVIDENCIA = {
    14: "01-inicio",                 # REQ-01 Arquitectura tecnologica
    15: "VS-02-estructura",          # REQ-02 Carpeta backend con FastAPI
    16: "TERM-03-entorno",           # REQ-03 Entorno FastAPI
    17: "VS-04-basedatos",           # REQ-04 Base de datos SQL
    18: "DB-05-usuarios",            # REQ-05 Tabla de usuarios
    19: "VS-06-modelos",             # REQ-06 Modelos y esquemas
    20: "VS-07-conexion",            # REQ-07 Conexion DB con FastAPI
    21: "14-admin-productos",        # REQ-08 Conexion FrontEnd/BackEnd/DB
    22: "09-registro-exitoso",       # REQ-09 Registro de clientes
    23: "05-login",                  # REQ-10 Inicio de sesion
    24: "POST-11-jwt",               # REQ-11 Autenticacion JWT
    25: "16-menu-usuario",           # REQ-12 Control de roles
    26: "POST-13-401",               # REQ-13 Proteccion de endpoints
    27: "swagger-01-general",        # REQ-14 Endpoints de la API
    28: "POST-15-registro",          # REQ-15 Metodos HTTP
    29: "13-admin-editar-usuario",   # REQ-16 Operaciones CRUD usuarios
    30: "12-admin-usuarios",         # REQ-17 Panel de administracion
    31: "17-empleado-productos",     # REQ-18 Panel de empleado
    32: "19-cliente-productos",      # REQ-19 Panel de cliente
    33: "11-navbar-admin",           # REQ-20 Usuario en el Navbar
    34: "08-registro-validaciones",  # REQ-21 Validaciones en tiempo real
    35: "TERM-22-hash",              # REQ-22 Seguridad de contrasenas
    36: "VS-23-entorno",             # REQ-23 Variables de entorno
    37: "02-quienes-somos",          # REQ-24 Componente WhatsApp
    38: "swagger-02-login",          # REQ-25 Documentacion FastAPI (Swagger)
    39: "POST-26-login",             # REQ-26 Pruebas con Postman
}

OBSERVACIONES = {
    14: "Frontend React + Vite consumiendo la API FastAPI en formato JSON (frontend/src/utils/api.js -> http://localhost:8000/api) y backend conectado a MySQL. La pagina de inicio lista los productos de GET /api/productos.",
    15: "Proyecto separado en frontend/ y backend-fastapi/, con subcarpeta app/ (main.py, config.py, database.py, models.py, schemas.py, security.py y routers/).",
    16: "Entorno virtual Python (.venv) con las dependencias instaladas desde requirements.txt: fastapi, uvicorn, sqlalchemy, pydantic, python-jose, bcrypt, pymysql, python-dotenv y email-validator.",
    17: "Base de datos relacional MySQL techpc_db definida en backend-fastapi/database.sql con roles, permisos, rol_permisos, usuarios, categorias_productos, productos y servicios.",
    18: "DESCRIBE usuarios muestra los campos exigidos (nombre, apellido, tipo_documento, numero_documento unico, direccion, telefono, correo unico, contrasena, rol_id y estado). El SELECT evidencia que la contrasena se guarda como hash bcrypt.",
    19: "app/models.py define los modelos SQLAlchemy y app/schemas.py los esquemas Pydantic, que validan tipos, campos obligatorios, longitudes, correo y contrasena.",
    20: "app/config.py construye la DATABASE_URL con las variables de entorno y app/database.py crea el engine de SQLAlchemy (pool_pre_ping) que consume la API.",
    21: "Comunicacion React -> FastAPI -> MySQL verificada: el panel de administracion muestra los productos y servicios que la API lee de la base de datos.",
    22: "Registro conectado a POST /api/usuarios/registro: validaciones en React y FastAPI, verificacion de no duplicados (correo y documento), hash bcrypt de la contrasena y respuesta JSON.",
    23: "Login conectado a POST /api/auth/login: FastAPI verifica las credenciales contra la base de datos y genera un JSON Web Token (JWT) con expiracion de 24 horas.",
    24: "El token se almacena en el navegador y se envia en la cabecera Authorization: Bearer <token> (frontend/src/utils/api.js). La API responde 200 al listar usuarios con token valido y 401 sin token.",
    25: "Roles administrador, empleado y cliente. El backend aplica la autorizacion definitiva con dependencias y React restringe los paneles segun el rol del usuario autenticado.",
    26: "Los endpoints protegidos con JWT responden 401 {'detail': 'Token no proporcionado.'} al llamarlos sin token. Ademas, abrir /admin sin sesion redirige al inicio.",
    27: "Endpoints REST en espanol para Usuarios, Productos y Servicios, agrupados por entidad en la documentacion: /api/auth, /api/usuarios, /api/productos y /api/servicios.",
    28: "Metodos implementados y probados contra la API en ejecucion: POST /api/usuarios/registro (201), GET /api/usuarios (200 con token, 401 sin token), PUT (200), PATCH (200) y DELETE (200).",
    29: "CRUD completo de usuarios: consultar, crear, editar (PUT /api/usuarios/{id}), cambiar estado activo/inactivo (PATCH /api/usuarios/{id}/estado) y eliminar (DELETE).",
    30: "Panel de administracion en React + Vite + Tailwind (ruta /admin) protegido por autenticacion y autorizacion. Gestiona usuarios, productos y servicios.",
    31: "Panel de empleado (/empleado) que solo muestra la gestion de productos y servicios. FastAPI valida el rol antes de permitir estas operaciones.",
    32: "Panel de cliente (/cliente) con productos, servicios, perfil y cambio de cuenta. El usuario se identifica con la informacion contenida en el JWT.",
    33: "El Navbar muestra el nombre del usuario autenticado y su menu (panel, cambiar cuenta, cerrar sesion). Se actualiza automaticamente al cerrar sesion.",
    34: "Validaciones en tiempo real en React (frontend/src/utils/validations.js) y re-validacion en FastAPI con Pydantic: obligatorios, longitudes, formato de correo, telefono y contrasena.",
    35: "El SELECT muestra la contrasena almacenada como hash $2b$10$... y bcrypt.checkpw('Admin123*', hash) devuelve True: no existe texto plano en la base de datos.",
    36: "Variables de entorno en backend-fastapi/.env (DB_HOST, DB_USER, DB_PASSWORD, DB_NAME, DB_PORT, PORT, JWT_SECRET) con su plantilla .env.example; el .env no se comparte.",
    37: "Se conserva el componente reutilizable WhatsAppButton.jsx con posicion fija y enlace https://wa.me/573044761403; funciona con independencia del backend.",
    38: "Swagger UI disponible en http://127.0.0.1:8000/docs. Se ejecuto POST /api/auth/login con 'Try it out' y respondio 200 OK con el token JWT.",
    39: "Coleccion Postman ejecutada contra la API: generacion de JWT, GET protegido con y sin token (200/401), registro de cliente (201) y operaciones PUT/PATCH/DELETE (200).",
}

IMG_W, IMG_H = 372, 252
ROW_H = 193


def main():
    wb = openpyxl.load_workbook(TEMPLATE)
    ws = wb.active

    # Datos del aprendiz
    ws["C5"] = "Anderson Florez Florez"
    ws["F5"] = "11/09/2026"
    ws["C6"] = "1020116685"
    ws["F6"] = "3406204"

    faltantes = []
    for row in range(14, 40):
        ws["D%d" % row] = ESTADO
        ws["F%d" % row] = OBSERVACIONES.get(row, "")

        nombre = EVIDENCIA.get(row)
        ruta = os.path.join(CAPTURAS, nombre + ".png") if nombre else None
        if not ruta or not os.path.exists(ruta):
            faltantes.append("fila %d (%s)" % (row, nombre))
            continue
        ws["E%d" % row] = None
        img = XLImage(ruta)
        img.width, img.height = IMG_W, IMG_H
        ws.add_image(img, "E%d" % row)
        ws.row_dimensions[row].height = ROW_H

    if faltantes:
        raise SystemExit("Faltan imagenes: " + ", ".join(faltantes))

    wb.save(SALIDA)
    print("Guardado:", SALIDA)
    print("Requerimientos con evidencia:", len(EVIDENCIA))


if __name__ == "__main__":
    main()
