"""Llena la lista de chequeo del QUINTO avance (React + Vite + FastAPI + SQL + IA).

- Marca el estado de los 25 requerimientos.
- Escribe las observaciones del aprendiz.
- Inserta la captura de evidencia de cada requerimiento (marcos de Windows 11).

Entrada : C:\\Users\\simon\\Downloads\\lista_Chequeo_Quinto_Avance_React_Fastapi.xlsx
Salida  : Lista_Chequeo_Quinto_Avance_REACT_FASTAPI_LLENADA.xlsx (raiz del proyecto)
"""
import os

import openpyxl
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

ROOT = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
PLANTILLA = r"C:\Users\simon\Downloads\lista_Chequeo_Quinto_Avance_React_Fastapi.xlsx"
EVIDENCIAS = os.path.join(ROOT, "docs", "capturas-quinto-avance")
SALIDA = os.path.join(ROOT, "Lista_Chequeo_Quinto_Avance_REACT_FASTAPI_LLENADA.xlsx")

GITHUB = "https://github.com/stronghoold/pagina_react_personal"

# fila -> (captura, estado)
EVIDENCIA = {
    15: ("08-admin-venta-registrada", "Cumplido"),
    16: ("09-admin-venta-detalle", "Cumplido"),
    17: ("10-admin-ventas-filtros", "Cumplido"),
    18: ("15-admin-reportes", "Cumplido"),
    19: ("E05-reporte-pdf", "Cumplido"),
    20: ("E06-reporte-excel", "Cumplido"),
    21: ("11-admin-ventas-facturada", "Cumplido"),
    22: ("12-admin-facturas", "Cumplido"),
    23: ("E09-factura-pdf", "Cumplido"),
    24: ("04-admin-dashboard", "Cumplido"),
    25: ("04b-admin-dashboard-graficos", "Cumplido"),
    26: ("30-empleado-dashboard", "Cumplido"),
    27: ("05-admin-dashboard-filtros", "Cumplido"),
    28: ("swagger-q5-01-general", "Cumplido"),
    29: ("E15-integracion-dashboard", "Cumplido"),
    30: ("18-admin-pqr", "Cumplido"),
    31: ("02-chatbot", "Cumplido"),
    32: ("E18-chatbot-ia", "Cumplido"),
    33: ("E19-api-key-segura", "Cumplido"),
    34: ("E20-despliegue", "En Proceso"),
    35: ("E21-base-datos-sql", "Cumplido"),
    36: ("E22-modelos-esquemas", "Cumplido"),
    37: ("E23-componentes-react", "Cumplido"),
    38: ("E24-seguridad", "Cumplido"),
    39: ("E25-postman", "Cumplido"),
}

OBSERVACIONES = {
    15: "Modulo de ventas en backend-fastapi/app/routers/ventas.py (POST /api/ventas) y en el componente VentasModal.jsx. La captura muestra la venta V-20260922-0001 registrada desde el panel: cliente, usuario operador, metodo de pago, estado, subtotal, impuestos y total, guardada de forma persistente en la tabla ventas de techpc_db.",
    16: "Detalle de la venta V-20260922-0001 con dos items: el producto RTX 5080 16GB ($ 3.500.000) y el servicio Actualizacion de Componentes ($ 60.000). Cada item se almacena en detalle_ventas con su cantidad, precio unitario y subtotal, relacionado con ventas, productos y servicios.",
    17: "Historial de ventas (VentasPanel.jsx + GET /api/ventas) con filtros por fecha inicial, fecha final, cliente, producto, servicio, estado, numero de venta y rango de valor ($ min - $ max). La captura corresponde a una consulta filtrada por estado y cliente.",
    18: "Reporte diario del 2026-09-22 obtenido con GET /api/ventas/reporte/diario: 1 venta, $ 4.236.400 de facturacion, $ 676.400 de impuestos y 2 unidades vendidas. La tabla incluye fecha, numero de venta, cliente, productos/servicios, cantidad, valor unitario, total y estado.",
    19: "Reporte diario exportado en PDF por FastAPI (GET /api/ventas/reporte/diario/pdf, generado con ReportLab en app/utils/reportes.py). El archivo reporte_ventas_2026-09-22.pdf incluye nombre del proyecto, ficha, fecha, tabla de ventas, totales y datos de generacion.",
    20: "Reporte diario exportado en Excel con GET /api/ventas/reporte/diario/excel (openpyxl). El archivo reporte_ventas_2026-09-22.xlsx organiza la informacion en 11 columnas (fecha, numero, cliente, producto/servicio, tipo, cantidad, precio unitario, descuento, subtotal, total y estado) listas para filtrar o procesar en Excel.",
    21: "Factura F-20260922-0001 generada desde la venta V-20260922-0001 con POST /api/facturas (app/routers/facturas.py). Queda relacionada 1:1 con la venta y almacena numero, fecha, cliente, items congelados en detalle_facturas, subtotal, impuestos, total y estado.",
    22: "Consulta de facturas (GET /api/facturas + FacturasPanel.jsx) con filtros por numero de factura, cliente, fechas y estado. El backend aplica el filtro del JWT, por lo que el cliente solo ve sus propias facturas.",
    23: "Descarga de la factura en PDF con GET /api/facturas/{id}/pdf. La captura muestra el documento factura_V-20260922-0001.pdf abierto: encabezado del proyecto, datos del cliente, detalle de items, subtotal, IVA del 19 % y total.",
    24: "Dashboard administrativo (GET /api/dashboard/admin + DashboardPanel.jsx) con Cards de usuarios, productos, servicios, ventas, facturacion total, impuestos y PQR pendientes. Todos los indicadores se calculan en el backend sobre la base de datos.",
    25: "Dashboard de ventas con grafico de barras de facturacion por dia, semana o mes, grafico lineal con la tendencia de los ultimos 14 dias, ventas por mes, productos mas vendidos y ventas por estado, ademas de Cards con las cifras filtradas.",
    26: "Dashboards diferenciados por rol: la captura corresponde al panel de empleado (indicadores comerciales y PQR) y el cliente tiene su propio resumen en /cliente. FastAPI protege cada endpoint con require_roles (administrador, empleado, cliente).",
    27: "Filtros de los dashboards por fecha inicial, fecha final, producto, servicio, estado, cliente y agrupacion (dia, semana o mes). Los indicadores, el grafico de barras y las listas se recalculan con GET /api/dashboard/ventas segun los filtros aplicados.",
    28: "Documentacion interactiva de FastAPI en http://127.0.0.1:8000/docs con los grupos nuevos del quinto avance: Ventas, Facturas, Dashboards, PQR y Chatbot (48 operaciones en total, visibles en Swagger UI).",
    29: "El Dashboard consume los endpoints de FastAPI: el codigo de app/routers/dashboard.py calcula las estadisticas con consultas SQLAlchemy (func.sum, func.count, group_by, filtros por fecha) sobre ventas, detalle_ventas, facturas y pqr. React no tiene datos quemados.",
    30: "Modulo de PQR (app/routers/pqr.py + PqrPanel.jsx): el cliente registra la solicitud y consulta su estado; el administrador y el empleado la gestionan con estado (pendiente, en proceso, respondida, cerrada) y respuesta. La captura muestra la gestion de una queja.",
    31: "Chatbot TechBot (components/Chatbot.jsx + POST /api/chat) integrado en toda la aplicacion: resuelve preguntas frecuentes y orienta sobre productos, precios, servicios, proceso de compra, facturas, garantias, envios y PQR.",
    32: "app/utils/ia.py integra desde FastAPI un servicio compatible con OpenAI: envia como contexto el catalogo real de productos y servicios y guarda cada conversacion en las tablas conversaciones y mensajes. Si la clave no esta configurada o el proveedor falla, responde el motor local de reglas (fuente \"local\") sin interrumpir la atencion.",
    33: "La clave del servicio de IA se lee con python-dotenv desde backend-fastapi/.env (IA_API_KEY, IA_BASE_URL, IA_MODEL, IA_TIMEOUT); nunca se escribe en el codigo ni se envia al navegador. El repositorio publica unicamente .env.example con la clave vacia y .env esta en .gitignore.",
    34: "EN PROCESO: el proyecto esta preparado para el despliegue con backend-fastapi/Procfile (uvicorn con $PORT), frontend/vercel.json y frontend/public/_redirects para el SPA, y las variables de entorno (DB_*, JWT_SECRET, CORS_ORIGINS, IA_*) para configurar en la plataforma. Falta publicar en Railway o Vercel y anexar la URL publica del frontend y de /docs.",
    35: "La base de datos techpc_db evoluciono con 7 tablas nuevas: ventas, detalle_ventas, facturas, detalle_facturas, pqr, conversaciones y mensajes, con claves foraneas hacia usuarios, productos y servicios. Scripts: backend-fastapi/database.sql y migracion_quinto_avance.sql.",
    36: "Nuevos modelos ORM con SQLAlchemy en app/models.py (Venta, DetalleVenta, Factura, DetalleFactura, Pqr, Conversacion y Mensaje) y esquemas Pydantic en app/schemas.py que validan tipos, campos obligatorios, cantidades, estados y longitudes de las peticiones y respuestas.",
    37: "Componentes React reutilizables nuevos: BarChart.jsx y LineChart.jsx (graficos SVG sin dependencias externas), StatCard.jsx (Cards), Chatbot.jsx, CheckoutModal.jsx y los paneles de la carpeta comercial (DashboardPanel, DashboardCliente, VentasPanel, VentasModal, FacturasPanel, ReportsPanel y PqrPanel), integrados con el diseno previo.",
    38: "Seguridad integral: los endpoints nuevos exigen JWT y rol (require_roles). La captura muestra GET /api/dashboard/admin sin token respondiendo 401 Unauthorized; ademas las contrasenas se almacenan con hash bcrypt y las credenciales viven en variables de entorno (.env).",
    39: "Coleccion Postman ejecutada contra la API en ejecucion: POST /api/ventas (201 Created), POST /api/facturas (201), GET /api/ventas y /api/facturas (200), reporte diario (200), PQR (201 y 200), chatbot (200) y las respuestas 401 de los endpoints protegidos sin token.",
}

IMG_W, IMG_H = 372, 252
ROW_H = 193


def main():
    wb = openpyxl.load_workbook(PLANTILLA)
    ws = wb.active

    # ── Datos del aprendiz ──
    ws["B5"] = "Anderson Florez Florez"
    ws["F5"] = "22/09/2026"
    ws["B6"] = "1020116685"
    ws["F6"] = "3406204"
    ws["F7"] = "03 / Ambiente 301"
    ws["B8"] = "React + Vite / FastAPI + IA"
    ws["F8"] = "Jhan Hader Muñoz"

    # La columna E (Estado) viene oculta en la plantilla: se muestra para que
    # los estados y los porcentajes sean visibles.
    ws.column_dimensions["E"].hidden = False
    ws.column_dimensions["E"].width = 15

    faltantes = []
    for fila, (nombre, estado) in EVIDENCIA.items():
        ws["E%d" % fila] = estado
        ws["E%d" % fila].alignment = Alignment(horizontal="center", vertical="center")
        ws["G%d" % fila] = OBSERVACIONES[fila]
        ws["G%d" % fila].alignment = Alignment(wrap_text=True, vertical="top")

        ruta = os.path.join(EVIDENCIAS, nombre + ".png")
        if not os.path.exists(ruta):
            faltantes.append("fila %d (%s)" % (fila, nombre))
            continue
        img = XLImage(ruta)
        img.width, img.height = IMG_W, IMG_H
        ws.add_image(img, "F%d" % fila)
        ws.row_dimensions[fila].height = ROW_H

    if faltantes:
        raise SystemExit("Faltan evidencias: " + ", ".join(faltantes))

    # ── URL del entregable ──
    ws["D40"] = GITHUB
    ws["D40"].font = Font(color="0563C1", underline="single", size=11)
    ws.merge_cells("D40:G40")

    # ── Ajustes finales del area del checklist ──
    ws.column_dimensions["F"].width = 54
    ws.column_dimensions["G"].width = 62
    ws["F14"].alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
    ws["G14"].alignment = Alignment(wrap_text=True, vertical="center")

    wb.save(SALIDA)
    cumplidos = sum(1 for _f, (n, e) in EVIDENCIA.items() if e == "Cumplido")
    print("Guardado:", SALIDA)
    print("Requerimientos con evidencia: %d" % len(EVIDENCIA))
    print("Cumplidos: %d | En proceso: %d" % (cumplidos, len(EVIDENCIA) - cumplidos))
    print("Columnas: Estado=%s Evidencia=%s Observaciones=%s" % ("E", "F", "G"))


if __name__ == "__main__":
    main()
