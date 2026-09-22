"""
Generación de documentos del quinto avance:

- Reporte diario de ventas en PDF  (requerimiento 5)
- Reporte diario de ventas en Excel (requerimiento 6)
- Factura de venta en PDF           (requerimiento 9)

Todos los documentos se construyen en memoria (BytesIO) y se devuelven como
descarga desde FastAPI, sin escribir archivos temporales en el servidor.
"""
from datetime import datetime
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import cm
from reportlab.platypus import (
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)

from ..config import IDENTIFICACION_PROYECTO, NOMBRE_PROYECTO

AZUL = colors.HexColor('#1E3A8A')
GRIS = colors.HexColor('#F1F5F9')


def _estilos():
    base = getSampleStyleSheet()
    return {
        'titulo': ParagraphStyle(
            'Titulo', parent=base['Title'], fontSize=16, textColor=AZUL, spaceAfter=2,
        ),
        'subtitulo': ParagraphStyle(
            'Subtitulo', parent=base['Normal'], fontSize=10, textColor=colors.HexColor('#475569'),
        ),
        'seccion': ParagraphStyle(
            'Seccion', parent=base['Heading3'], fontSize=11, textColor=AZUL, spaceBefore=10,
        ),
        'normal': ParagraphStyle('Normal2', parent=base['Normal'], fontSize=8.5),
        'celda': ParagraphStyle('Celda', parent=base['Normal'], fontSize=8),
    }


def _encabezado_documento(e, titulo: str, subtitulo: str):
    """Bloque de encabezado común a reportes y facturas."""
    e.append(Paragraph(NOMBRE_PROYECTO, _estilos()['titulo']))
    e.append(Paragraph(IDENTIFICACION_PROYECTO, _estilos()['subtitulo']))
    e.append(Spacer(1, 6))
    e.append(Paragraph(titulo, _estilos()['seccion']))
    e.append(Paragraph(subtitulo, _estilos()['subtitulo']))
    e.append(Spacer(1, 8))


def _tabla(data, anchos=None, alinear_derecha_desde: int | None = None, encabezado=True):
    tabla = Table(data, colWidths=anchos, repeatRows=1 if encabezado else 0)
    comandos = [
        ('GRID', (0, 0), (-1, -1), 0.4, colors.HexColor('#CBD5E1')),
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
        ('FONTSIZE', (0, 0), (-1, -1), 8),
        ('LEFTPADDING', (0, 0), (-1, -1), 4),
        ('RIGHTPADDING', (0, 0), (-1, -1), 4),
    ]
    if encabezado:
        comandos += [
            ('BACKGROUND', (0, 0), (-1, 0), AZUL),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
            ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ]
        if len(data) > 1:
            comandos.append(('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, GRIS]))
    if alinear_derecha_desde is not None:
        comandos.append(('ALIGN', (alinear_derecha_desde, 0), (-1, -1), 'RIGHT'))
    tabla.setStyle(TableStyle(comandos))
    return tabla


def _moneda(valor) -> str:
    return f'$ {float(valor or 0):,.0f}'.replace(',', '.')


# ════════════════════════════════════════════
# Reporte diario de ventas
# ════════════════════════════════════════════

def reporte_ventas_pdf(reporte: dict, generado_por: str) -> BytesIO:
    """Construye el PDF del reporte diario de ventas."""
    buffer = BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=letter,
        leftMargin=1.6 * cm,
        rightMargin=1.6 * cm,
        topMargin=1.6 * cm,
        bottomMargin=1.6 * cm,
        title=f"Reporte diario de ventas {reporte['fecha']}",
    )
    estilos = _estilos()
    e = []
    _encabezado_documento(
        e,
        'Reporte diario de ventas',
        f"Fecha del reporte: {reporte['fecha']}",
    )

    # ── Resumen ──
    resumen = reporte['resumen']
    e.append(Paragraph('Resumen del día', estilos['seccion']))
    filas_resumen = [
        ['Ventas registradas', 'Facturación total', 'Impuestos', 'Unidades vendidas'],
        [
            str(resumen['total_ventas']),
            _moneda(resumen['facturacion']),
            _moneda(resumen['impuestos']),
            str(resumen['unidades']),
        ],
    ]
    e.append(_tabla(filas_resumen, anchos=[4.3 * cm] * 4))
    e.append(Spacer(1, 10))

    # ── Detalle de ventas ──
    e.append(Paragraph('Detalle de ventas', estilos['seccion']))
    data = [['N° venta', 'Cliente', 'Producto / Servicio', 'Cant.', 'Valor unit.', 'Total venta', 'Estado']]
    for venta in reporte['ventas']:
        items = venta['items'] or [{'descripcion': '—', 'cantidad': 0, 'precio_unitario': 0}]
        primera = True
        for item in items:
            data.append([
                venta['numero_venta'] if primera else '',
                (venta['cliente_nombre'] or '') if primera else '',
                Paragraph(item['descripcion'], estilos['celda']),
                str(item['cantidad']),
                _moneda(item['precio_unitario']),
                _moneda(venta['total']) if primera else '',
                venta['estado'] if primera else '',
            ])
            primera = False

    if len(data) == 1:
        data.append(['—', 'Sin ventas registradas en esta fecha', '', '', '', '', ''])

    e.append(_tabla(
        data,
        anchos=[2.4 * cm, 3.2 * cm, 5.0 * cm, 1.2 * cm, 2.2 * cm, 2.2 * cm, 1.8 * cm],
        encabezado=True,
    ))
    e.append(Spacer(1, 12))

    # ── Información de generación ──
    e.append(Paragraph(
        f"Documento generado automáticamente por {NOMBRE_PROYECTO} el "
        f"{datetime.now().strftime('%d/%m/%Y a las %H:%M:%S')} "
        f"por el usuario {generado_por}.",
        estilos['normal'],
    ))
    doc.build(e)
    buffer.seek(0)
    return buffer


def reporte_ventas_excel(reporte: dict, generado_por: str) -> BytesIO:
    """Construye el archivo Excel (.xlsx) del reporte diario de ventas."""
    wb = Workbook()
    ws = wb.active
    ws.title = 'Reporte diario'

    titulo_font = Font(bold=True, size=14, color='1E3A8A')
    encabezado_font = Font(bold=True, color='FFFFFF')
    encabezado_fill = PatternFill('solid', fgColor='1E3A8A')
    borde = Border(*[Side(style='thin', color='CBD5E1')] * 4)

    ws['A1'] = NOMBRE_PROYECTO
    ws['A1'].font = titulo_font
    ws['A2'] = IDENTIFICACION_PROYECTO
    ws['A3'] = f"Reporte diario de ventas — {reporte['fecha']}"
    ws['A3'].font = Font(bold=True, size=11)
    ws['A4'] = f"Generado el {datetime.now().strftime('%d/%m/%Y %H:%M:%S')} por {generado_por}"
    ws['A4'].font = Font(italic=True, size=9, color='64748B')

    encabezados = [
        'Fecha', 'N° venta', 'Cliente', 'Producto / Servicio', 'Tipo', 'Cantidad',
        'Precio unitario', 'Descuento', 'Subtotal item', 'Total venta', 'Estado',
    ]
    fila_inicio = 6
    for col, texto in enumerate(encabezados, start=1):
        celda = ws.cell(row=fila_inicio, column=col, value=texto)
        celda.font = encabezado_font
        celda.fill = encabezado_fill
        celda.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        celda.border = borde

    fila = fila_inicio + 1
    for venta in reporte['ventas']:
        items = venta['items'] or [{
            'descripcion': '—', 'tipo': '', 'cantidad': 0,
            'precio_unitario': 0, 'descuento': 0, 'subtotal': 0,
        }]
        for item in items:
            valores = [
                str(venta['fecha'])[:10],
                venta['numero_venta'],
                venta['cliente_nombre'],
                item['descripcion'],
                item.get('tipo', ''),
                item['cantidad'],
                float(item['precio_unitario']),
                float(item.get('descuento', 0)),
                float(item.get('subtotal', 0)),
                float(venta['total']),
                venta['estado'],
            ]
            for col, valor in enumerate(valores, start=1):
                celda = ws.cell(row=fila, column=col, value=valor)
                celda.border = borde
                if col in {7, 8, 9, 10}:
                    celda.number_format = '"$" #,##0'
            fila += 1

    if fila == fila_inicio + 1:
        ws.cell(row=fila, column=1, value='Sin ventas registradas en esta fecha')
        fila += 1

    # ── Totales ──
    fila += 1
    resumen = reporte['resumen']
    totales = [
        ('Total de facturación', resumen['facturacion']),
        ('Impuestos', resumen['impuestos']),
        ('Descuentos', resumen['descuentos']),
        ('Ventas registradas', resumen['total_ventas']),
        ('Unidades vendidas', resumen['unidades']),
    ]
    for etiqueta, valor in totales:
        ws.cell(row=fila, column=1, value=etiqueta).font = Font(bold=True)
        celda = ws.cell(row=fila, column=2, value=valor)
        if etiqueta not in {'Ventas registradas', 'Unidades vendidas'}:
            celda.number_format = '"$" #,##0'
        fila += 1

    anchos = [11, 14, 26, 40, 11, 9, 15, 12, 14, 14, 11]
    for i, ancho in enumerate(anchos, start=1):
        ws.column_dimensions[get_column_letter(i)].width = ancho
    ws.freeze_panes = ws.cell(row=fila_inicio + 1, column=1)

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer


# ════════════════════════════════════════════
# Factura de venta
# ════════════════════════════════════════════

def factura_pdf(factura: dict) -> BytesIO:
    """Construye el PDF de una factura de venta."""
    buffer = BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=letter,
        leftMargin=1.6 * cm,
        rightMargin=1.6 * cm,
        topMargin=1.6 * cm,
        bottomMargin=1.6 * cm,
        title=f"Factura {factura['numero_factura']}",
    )
    estilos = _estilos()
    e = []
    _encabezado_documento(
        e,
        f"Factura de venta N° {factura['numero_factura']}",
        f"Venta asociada: {factura.get('numero_venta') or '—'}  |  "
        f"Fecha: {str(factura['fecha'])[:19].replace('T', ' ')}",
    )

    # ── Datos del cliente ──
    e.append(Paragraph('Datos del cliente', estilos['seccion']))
    e.append(_tabla(
        [
            ['Nombre', 'Documento', 'Correo'],
            [
                factura.get('cliente_nombre') or '—',
                factura.get('cliente_documento') or '—',
                factura.get('cliente_correo') or '—',
            ],
        ],
        anchos=[6.0 * cm, 4.5 * cm, 7.0 * cm],
    ))
    e.append(Spacer(1, 10))

    # ── Detalle ──
    e.append(Paragraph('Productos y servicios', estilos['seccion']))
    data = [['Descripción', 'Cantidad', 'Precio unitario', 'Descuento', 'Subtotal']]
    for item in factura['items']:
        data.append([
            Paragraph(item['descripcion'], estilos['celda']),
            str(item['cantidad']),
            _moneda(item['precio_unitario']),
            _moneda(item['descuento']),
            _moneda(item['subtotal']),
        ])
    e.append(_tabla(data, anchos=[7.5 * cm, 2.0 * cm, 3.0 * cm, 2.5 * cm, 2.5 * cm], alinear_derecha_desde=1))
    e.append(Spacer(1, 10))

    # ── Totales ──
    e.append(Paragraph('Resumen de la factura', estilos['seccion']))
    e.append(_tabla(
        [['Subtotal', 'Descuento', 'Impuestos', 'Total', 'Estado', 'Método de pago'],
         [
             _moneda(factura['subtotal']),
             _moneda(factura['descuento']),
             _moneda(factura['impuestos']),
             _moneda(factura['total']),
             factura['estado'],
             factura.get('metodo_pago') or '—',
         ]],
        anchos=[2.7 * cm] * 6,
    ))
    e.append(Spacer(1, 14))
    e.append(Paragraph(
        f"Documento generado automáticamente por {NOMBRE_PROYECTO} el "
        f"{datetime.now().strftime('%d/%m/%Y a las %H:%M:%S')}.",
        estilos['normal'],
    ))
    doc.build(e)
    buffer.seek(0)
    return buffer
